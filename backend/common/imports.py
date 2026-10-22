"""CSV / XLSX 导入文件解析。"""
import csv
import io

from openpyxl import load_workbook

from common.exceptions import BusinessError


def read_table(uploaded_file, max_size_mb: int = 5) -> list[list[str]]:
    """读取上传表格，返回按行组织的字符串二维表。"""
    if uploaded_file.size > max_size_mb * 1024 * 1024:
        raise BusinessError(400, f"导入文件不能超过 {max_size_mb} MB")

    name = (uploaded_file.name or "").lower()
    try:
        if name.endswith(".csv"):
            raw = uploaded_file.read()
            text = raw.decode("utf-8-sig")
            return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]
        if name.endswith(".xlsx"):
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
            sheet = workbook.active
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell).strip() for cell in row])
            workbook.close()
            return rows
    except BusinessError:
        raise
    except Exception as exc:
        raise BusinessError(400, "导入文件解析失败，请检查文件格式") from exc

    raise BusinessError(400, "仅支持 CSV 或 XLSX 文件")


def header_map(header_row: list[str]) -> dict[str, int]:
    return {name.strip(): index for index, name in enumerate(header_row) if name.strip()}


def cell(row: list[str], mapping: dict[str, int], name: str) -> str:
    index = mapping.get(name)
    if index is None or index >= len(row):
        return ""
    return row[index].strip()
